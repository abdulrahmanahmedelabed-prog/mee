"""Report definitions (BioTime 'Reports' module) and CSV / Excel export."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, time, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import models as m
from .adms.protocol import VERIFY_TYPES
from .engine import Engine, summarize

STATUS_AR = {
    "present": "حاضر", "late": "متأخر", "early": "خروج مبكر", "late_early": "متأخر وخروج مبكر",
    "absent": "غائب", "leave": "إجازة", "holiday": "عطلة رسمية", "off": "راحة",
    "incomplete": "بصمة ناقصة", "unscheduled": "بدون جدول", "pending": "لم يحضر بعد", "": "",
}
STATUS_EN = {
    "present": "Present", "late": "Late", "early": "Early leave", "late_early": "Late & early",
    "absent": "Absent", "leave": "Leave", "holiday": "Holiday", "off": "Day off",
    "incomplete": "Missed punch", "unscheduled": "Unscheduled", "pending": "Not yet", "": "",
}
STATE_AR = {0: "دخول", 1: "خروج", 2: "خروج استراحة", 3: "عودة استراحة", 4: "دخول إضافي", 5: "خروج إضافي", 255: "-"}
STATE_EN = {0: "Check-In", 1: "Check-Out", 2: "Break-Out", 3: "Break-In", 4: "OT-In", 5: "OT-Out", 255: "-"}
VERIFY_AR = {"password": "كلمة مرور", "fingerprint": "بصمة إصبع", "card": "بطاقة", "face": "وجه",
             "palm": "كف", "finger_vein": "وريد الإصبع", "other": "أخرى"}
MONTH_SYMBOL = {"present": "✓", "late": "L", "early": "E", "late_early": "LE", "absent": "A",
                "leave": "V", "holiday": "H", "off": "-", "incomplete": "!", "unscheduled": "•", "pending": "", "": ""}

# key -> (Arabic title, English title, kind)
REPORTS = {
    "transactions": ("سجل الحركات", "Transactions", "punch"),
    "first_last": ("أول وآخر بصمة", "First & Last", "day"),
    "time_card": ("بطاقة الدوام", "Time Card", "day"),
    "daily": ("الحضور اليومي", "Daily Attendance", "day"),
    "late": ("تقرير التأخير", "Late Report", "day"),
    "early": ("تقرير الخروج المبكر", "Early Leave", "day"),
    "absent": ("تقرير الغياب", "Absence Report", "day"),
    "overtime": ("تقرير العمل الإضافي", "Overtime Report", "day"),
    "exception": ("البصمات الناقصة", "Missed Punch", "day"),
    "leave": ("تقرير الإجازات", "Leave Report", "leave"),
    "summary": ("الملخص الشهري", "Monthly Summary", "summary"),
    "monthly_status": ("كشف الحضور الشهري", "Monthly Status", "matrix"),
    "department": ("ملخص الأقسام", "Department Summary", "dept"),
}


def hm(mins) -> str:
    try:
        mins = int(mins or 0)
    except (TypeError, ValueError):
        return str(mins)
    if not mins:
        return ""
    return f"{mins // 60}:{mins % 60:02d}"


def _t(d: str) -> str:
    return d[11:16] if d else ""


def _cols(lang: str, spec: list[tuple[str, str, str]]) -> list[dict]:
    return [{"key": k, "label": ar if lang == "ar" else en} for k, ar, en in spec]


def build(db: Session, key: str, start: date, end: date, *, lang: str = "ar",
          employee_ids=None, department_ids=None, device_sn: str | None = None) -> dict:
    """Returns {"title", "columns": [{key,label}], "rows": [dict], "kind"}."""
    if key not in REPORTS:
        raise KeyError(key)
    title_ar, title_en, kind = REPORTS[key]
    status_names = STATUS_AR if lang == "ar" else STATUS_EN
    base = [("emp_code", "الرقم", "ID"), ("name", "الاسم", "Name"), ("department", "القسم", "Department")]
    out = {"title": title_ar if lang == "ar" else title_en, "kind": kind, "key": key,
           "start": start.isoformat(), "end": end.isoformat()}

    if kind == "punch":
        q = select(m.Transaction, m.Employee).outerjoin(m.Employee, m.Employee.id == m.Transaction.employee_id).where(
            m.Transaction.punch_time >= datetime.combine(start, time.min),
            m.Transaction.punch_time < datetime.combine(end + timedelta(days=1), time.min))
        if employee_ids:
            q = q.where(m.Transaction.employee_id.in_(employee_ids))
        if department_ids:
            q = q.where(m.Employee.department_id.in_(department_ids))
        if device_sn:
            q = q.where(m.Transaction.device_sn == device_sn)
        aliases = dict(db.execute(select(m.Device.sn, m.Device.alias)).all())
        rows = []
        for t, e in db.execute(q.order_by(m.Transaction.punch_time)).all():
            v = VERIFY_TYPES.get(t.verify_type, "other")
            rows.append({
                "emp_code": t.emp_code, "name": e.full_name if e else "",
                "department": e.department.name if e and e.department else "",
                "date": t.punch_time.strftime("%Y-%m-%d"), "time": t.punch_time.strftime("%H:%M:%S"),
                "state": (STATE_AR if lang == "ar" else STATE_EN).get(t.punch_state, str(t.punch_state)),
                "verify": VERIFY_AR.get(v, v) if lang == "ar" else v.replace("_", " ").title(),
                "device": aliases.get(t.device_sn, t.device_sn) or ("يدوي" if lang == "ar" else "Manual"),
                "temperature": t.temperature if t.temperature is not None else "",
            })
        out["columns"] = _cols(lang, base + [("date", "التاريخ", "Date"), ("time", "الوقت", "Time"),
                                             ("state", "الحالة", "State"), ("verify", "طريقة التحقق", "Verify"),
                                             ("device", "الجهاز", "Device"), ("temperature", "الحرارة", "Temp.")])
        out["rows"] = rows
        return out

    if kind == "leave":
        q = select(m.Leave, m.Employee, m.LeaveType).join(m.Employee, m.Employee.id == m.Leave.employee_id).join(
            m.LeaveType, m.LeaveType.id == m.Leave.leave_type_id).where(
            m.Leave.start_time < datetime.combine(end + timedelta(days=1), time.min),
            m.Leave.end_time > datetime.combine(start, time.min))
        if employee_ids:
            q = q.where(m.Leave.employee_id.in_(employee_ids))
        if department_ids:
            q = q.where(m.Employee.department_id.in_(department_ids))
        st = {"approved": "معتمدة", "pending": "بانتظار الموافقة", "rejected": "مرفوضة"} if lang == "ar" else {}
        rows = [{
            "emp_code": e.emp_code, "name": e.full_name, "department": e.department.name if e.department else "",
            "type": lt.name, "start": lv.start_time.strftime("%Y-%m-%d %H:%M"),
            "end": lv.end_time.strftime("%Y-%m-%d %H:%M"),
            "days": round((lv.end_time - lv.start_time).total_seconds() / 86400, 2),
            "status": st.get(lv.status, lv.status), "reason": lv.reason,
        } for lv, e, lt in db.execute(q.order_by(m.Leave.start_time)).all()]
        out["columns"] = _cols(lang, base + [("type", "النوع", "Type"), ("start", "من", "From"),
                                             ("end", "إلى", "To"), ("days", "أيام", "Days"),
                                             ("status", "الحالة", "Status"), ("reason", "السبب", "Reason")])
        out["rows"] = rows
        return out

    days = Engine(db, start, end, employee_ids=employee_ids, department_ids=department_ids).run()

    if kind == "summary":
        rows = summarize(days)
        for r in rows:
            r["leave_detail"] = ", ".join(f"{k}:{v}" for k, v in sorted(r.pop("leave_by_code").items()))
            for k in ("required", "worked", "late", "early", "ot", "leave"):
                r[k + "_hm"] = hm(r[k])
        out["columns"] = _cols(lang, base + [
            ("scheduled_days", "أيام الدوام", "Work days"), ("present_days", "أيام الحضور", "Present"),
            ("absent_days", "أيام الغياب", "Absent"), ("late_count", "مرات التأخير", "Late times"),
            ("late_hm", "مدة التأخير", "Late"), ("early_count", "مرات الخروج المبكر", "Early times"),
            ("early_hm", "مدة الخروج المبكر", "Early"), ("missed_punch", "بصمات ناقصة", "Missed"),
            ("required_hm", "الساعات المطلوبة", "Required"), ("worked_hm", "ساعات العمل", "Worked"),
            ("ot_hm", "العمل الإضافي", "Overtime"), ("leave_days", "أيام الإجازة", "Leave days"),
            ("leave_detail", "تفصيل الإجازات", "Leave detail"), ("holidays", "العطل", "Holidays"),
            ("off_days", "أيام الراحة", "Days off")])
        out["rows"] = rows
        return out

    if kind == "matrix":
        dates = []
        d = start
        while d <= end:
            dates.append(d)
            d += timedelta(days=1)
        by_emp: dict[int, dict] = {}
        for r in days:
            row = by_emp.setdefault(r.employee_id, {"emp_code": r.emp_code, "name": r.name,
                                                     "department": r.department, "_p": 0, "_a": 0, "_l": 0})
            sym = MONTH_SYMBOL.get(r.status, "")
            if r.status == "leave" and r.leave_codes:
                sym = r.leave_codes[0]
            row[r.att_date.isoformat()] = sym
            row["_p"] += 1 if r.status in ("present", "late", "early", "late_early", "incomplete") else 0
            row["_a"] += 1 if r.status == "absent" else 0
            row["_l"] += 1 if r.late else 0
        cols = base + [(d.isoformat(), f"{d.day}", f"{d.day}") for d in dates]
        cols += [("_p", "حضور", "P"), ("_a", "غياب", "A"), ("_l", "تأخير", "L")]
        out["columns"] = _cols(lang, cols)
        out["rows"] = list(by_emp.values())
        out["legend"] = {v: (STATUS_AR if lang == "ar" else STATUS_EN)[k] for k, v in MONTH_SYMBOL.items() if v}
        return out

    if kind == "dept":
        per: dict[str, dict] = {}
        for s in summarize(days):
            d = per.setdefault(s["department"] or "-", {"department": s["department"] or "-", "employees": 0,
                                                         "present_days": 0.0, "absent_days": 0, "late_count": 0,
                                                         "late": 0, "early": 0, "worked": 0, "ot": 0, "leave_days": 0})
            d["employees"] += 1
            for k in ("present_days", "absent_days", "late_count", "late", "early", "worked", "ot", "leave_days"):
                d[k] += s[k]
        rows = []
        for d in per.values():
            d["present_days"] = round(d["present_days"], 2)
            for k in ("late", "early", "worked", "ot"):
                d[k + "_hm"] = hm(d[k])
            rows.append(d)
        out["columns"] = _cols(lang, [
            ("department", "القسم", "Department"), ("employees", "الموظفون", "Employees"),
            ("present_days", "أيام الحضور", "Present days"), ("absent_days", "أيام الغياب", "Absent days"),
            ("late_count", "مرات التأخير", "Late times"), ("late_hm", "مدة التأخير", "Late"),
            ("early_hm", "الخروج المبكر", "Early"), ("worked_hm", "ساعات العمل", "Worked"),
            ("ot_hm", "العمل الإضافي", "Overtime"), ("leave_days", "أيام الإجازة", "Leave days")])
        out["rows"] = rows
        return out

    # ---- day based reports
    filt = {
        "late": lambda r: r.late > 0,
        "early": lambda r: r.early > 0,
        "absent": lambda r: r.status == "absent",
        "overtime": lambda r: r.ot > 0,
        "exception": lambda r: any(x in ("missed_in", "missed_out") for x in r.exceptions),
        "first_last": lambda r: bool(r.punches),
        "time_card": lambda r: True,
        "daily": lambda r: True,
    }[key]
    rows = []
    for r in days:
        if not filt(r):
            continue
        d = r.as_dict()
        d["clock_in"], d["clock_out"] = _t(d["clock_in"]), _t(d["clock_out"])
        d["sched"] = f"{_t(d['sched_in'])}-{_t(d['sched_out'])}" if d["sched_in"] else ""
        d["status_label"] = status_names.get(r.status, r.status)
        if r.holiday:
            d["status_label"] += f" ({r.holiday})"
        if r.leave_codes:
            d["status_label"] += " [" + ",".join(r.leave_codes) + "]"
        d["punch_list"] = " ".join(d["punches"])
        d["first"] = d["punches"][0] if d["punches"] else ""
        d["last"] = d["punches"][-1] if len(d["punches"]) > 1 else ""
        for k in ("required", "worked", "late", "early", "absent", "ot", "leave"):
            d[k + "_hm"] = hm(d[k])
        missed = {"missed_in": "بدون دخول" if lang == "ar" else "No check-in",
                  "missed_out": "بدون خروج" if lang == "ar" else "No check-out"}
        d["missed"] = "، ".join(missed[x] for x in r.exceptions if x in missed)
        rows.append(d)
    day_cols = [("date", "التاريخ", "Date")] + base
    spec = {
        "first_last": day_cols + [("first", "أول بصمة", "First punch"), ("last", "آخر بصمة", "Last punch"),
                                  ("punch_list", "كل البصمات", "All punches")],
        "time_card": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"),
                                 ("punch_list", "البصمات", "Punches"), ("worked_hm", "ساعات العمل", "Worked"),
                                 ("status_label", "الحالة", "Status")],
        "daily": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"),
                             ("clock_in", "الدخول", "Clock in"), ("clock_out", "الخروج", "Clock out"),
                             ("late_hm", "تأخير", "Late"), ("early_hm", "خروج مبكر", "Early"),
                             ("absent_hm", "غياب", "Absent"), ("worked_hm", "عمل", "Worked"),
                             ("ot_hm", "إضافي", "OT"), ("leave_hm", "إجازة", "Leave"),
                             ("status_label", "الحالة", "Status")],
        "late": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"),
                            ("late_hm", "مدة التأخير", "Late")],
        "early": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_out", "الخروج", "Clock out"),
                             ("early_hm", "مدة الخروج المبكر", "Early")],
        "absent": day_cols + [("timetable", "الدوام", "Timetable"), ("sched", "الوقت المقرر", "Schedule"),
                              ("absent_hm", "مدة الغياب", "Absent")],
        "overtime": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"),
                                ("clock_out", "الخروج", "Clock out"), ("ot_hm", "العمل الإضافي", "Overtime"),
                                ("status_label", "الحالة", "Status")],
        "exception": day_cols + [("sched", "الوقت المقرر", "Schedule"), ("clock_in", "الدخول", "Clock in"),
                                 ("clock_out", "الخروج", "Clock out"), ("missed", "النقص", "Missing")],
    }[key]
    out["columns"] = _cols(lang, spec)
    out["rows"] = rows
    return out


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

def to_csv(report: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in report["columns"]])
    for r in report["rows"]:
        w.writerow([r.get(c["key"], "") for c in report["columns"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel shows Arabic correctly


def to_xlsx(report: dict, company: str = "", rtl: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:30]
    ws.sheet_view.rightToLeft = rtl
    ncol = max(1, len(report["columns"]))
    ws.cell(row=1, column=1, value=company).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"{report['title']}   {report.get('start', '')} → {report.get('end', '')}").font = Font(bold=True, size=12)
    head_fill = PatternFill("solid", fgColor="1F6FB2")
    thin = Side(style="thin", color="BFC9D6")
    for i, c in enumerate(report["columns"], start=1):
        cell = ws.cell(row=4, column=i, value=c["label"])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for ri, r in enumerate(report["rows"], start=5):
        for ci, c in enumerate(report["columns"], start=1):
            v = r.get(c["key"], "")
            if isinstance(v, (list, dict)):
                v = str(v)
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for i, c in enumerate(report["columns"], start=1):
        width = max([len(str(c["label"]))] + [len(str(r.get(c["key"], ""))) for r in report["rows"][:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(5, width + 2))
    ws.freeze_panes = "A5"
    if ncol > 12:
        ws.page_setup.orientation = "landscape"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
